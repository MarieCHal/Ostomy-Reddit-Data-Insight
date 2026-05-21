"""Shared Excel merge, distribution, and formatting for theme/corpus reports."""

from __future__ import annotations

import json
from typing import Any

import pandas as pd


def category_order(meta_themes: dict[str, Any]) -> list[dict[str, str]]:
    raw = meta_themes.get("categories")
    if isinstance(raw, list) and raw:
        out: list[dict[str, str]] = []
        for item in raw:
            if isinstance(item, dict) and item.get("short_name"):
                out.append(
                    {
                        "id": str(item.get("id", "")),
                        "short_name": str(item["short_name"]),
                        "display_name": str(item.get("display_name", "")),
                    }
                )
        if out:
            return out
    return []


def infer_category_order(
    themes_by_id: dict[str, dict[str, Any]],
) -> list[dict[str, str]]:
    for row in themes_by_id.values():
        scores = row.get("theme_scores")
        if isinstance(scores, dict) and scores:
            return [{"id": "", "short_name": k, "display_name": ""} for k in scores]
    return []


def merge_posts_and_themes(
    posts: list[dict[str, Any]],
    themes_by_id: dict[str, dict[str, Any]],
    categories: list[dict[str, str]],
    threshold: float | None,
) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []
    short_names = [c["short_name"] for c in categories]

    for row in posts:
        pid = str(row.get("id", ""))
        t = themes_by_id.get(pid, {})
        title = (row.get("title") or "") or ""
        body = (row.get("selftext") or "") or ""
        text = f"{title}\n{body}".strip() if title and body else (title or body).strip()
        scores = t.get("theme_scores")
        if not isinstance(scores, dict):
            scores = {}

        theme_labels = t.get("theme_labels")
        if not isinstance(theme_labels, list):
            theme_labels = []

        scores_json = json.dumps(scores, ensure_ascii=False) if scores else ""

        record: dict[str, Any] = {
            "id": row.get("id"),
            "created_utc": row.get("created_utc"),
            "title": title,
            "selftext": body,
            "text_for_review": text,
            "theme_labels": "\n".join(str(x) for x in theme_labels),
            "theme_label": t.get("theme_label"),
            "theme_score": t.get("theme_score"),
            "theme_scores_json": scores_json,
            "theme_classifier_note": t.get("theme_classifier_note", ""),
        }

        for short_name in short_names:
            score = float(scores.get(short_name, 0.0))
            record[f"score_{short_name}"] = score
            if theme_labels:
                record[f"label_{short_name}"] = int(short_name in theme_labels)
            elif threshold is not None:
                record[f"label_{short_name}"] = int(score >= threshold)
            else:
                record[f"label_{short_name}"] = 0

        merged.append(record)
    return merged


def theme_distribution(
    rows: list[dict[str, Any]],
    categories: list[dict[str, str]],
) -> pd.DataFrame:
    total = len(rows) or 1
    out: list[dict[str, Any]] = []
    for cat in categories:
        short_name = cat["short_name"]
        col = f"label_{short_name}"
        count = sum(1 for r in rows if r.get(col) == 1)
        out.append(
            {
                "category_id": cat.get("id", ""),
                "short_name": short_name,
                "display_name": cat.get("display_name", ""),
                "count": count,
                "percent": round(100.0 * count / total, 2),
            }
        )
    out.sort(key=lambda x: (-x["count"], x["short_name"]))
    return pd.DataFrame(out)


def posts_column_order(
    df: pd.DataFrame,
    categories: list[dict[str, str]],
) -> list[str]:
    short_names = [c["short_name"] for c in categories]
    front = [
        "id",
        "created_utc",
        "title",
        "selftext",
        "text_for_review",
        "theme_labels",
        "theme_label",
        "theme_score",
    ]
    label_cols = [f"label_{s}" for s in short_names if f"label_{s}" in df.columns]
    score_cols = [f"score_{s}" for s in short_names if f"score_{s}" in df.columns]
    tail = ["theme_classifier_note", "theme_scores_json"]
    ordered: list[str] = []
    for col in front + label_cols + score_cols + tail:
        if col in df.columns and col not in ordered:
            ordered.append(col)
    for col in df.columns:
        if col not in ordered:
            ordered.append(col)
    return ordered


def _estimate_row_height(
    text: str, col_width: int, min_h: float = 30, max_h: float = 400
) -> float:
    if not text:
        return min_h
    chars_per_line = max(col_width * 1.1, 20)
    lines = max(1, len(str(text)) / chars_per_line)
    explicit = str(text).count("\n") + 1
    lines = max(lines, explicit)
    return min(max_h, max(min_h, lines * 15))


def format_posts_worksheet(ws: Any, column_names: list[str]) -> None:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wrap_top = Alignment(wrap_text=True, vertical="top")
    header_font = Font(bold=True)
    header_align = Alignment(wrap_text=True, vertical="center")

    col_widths: dict[str, float] = {
        "id": 14,
        "created_utc": 20,
        "title": 36,
        "selftext": 48,
        "text_for_review": 56,
        "theme_labels": 28,
        "theme_label": 28,
        "theme_score": 12,
        "theme_classifier_note": 16,
        "theme_scores_json": 24,
        "sentiment_label": 14,
        "sentiment_polarity_index": 16,
        "sentiment_score": 12,
        "sentiment_scores_json": 22,
        "sentiment_classifier_note": 16,
    }

    review_cols = {
        "title",
        "selftext",
        "text_for_review",
        "theme_labels",
        "theme_scores_json",
        "sentiment_scores_json",
    }

    for col_idx, name in enumerate(column_names, start=1):
        letter = get_column_letter(col_idx)
        if name.startswith("label_"):
            width = 10
        elif name.startswith("score_"):
            width = 11
        else:
            width = col_widths.get(name, 14)
        ws.column_dimensions[letter].width = width

        header = ws.cell(row=1, column=col_idx)
        header.font = header_font
        header.alignment = header_align

        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if name in review_cols or name.startswith("label_"):
                cell.alignment = wrap_top

    ws.freeze_panes = "D2"

    text_idx = (
        column_names.index("text_for_review") + 1
        if "text_for_review" in column_names
        else None
    )
    for row_idx in range(2, ws.max_row + 1):
        if text_idx:
            text = ws.cell(row=row_idx, column=text_idx).value or ""
            ws.row_dimensions[row_idx].height = _estimate_row_height(
                str(text), col_widths.get("text_for_review", 56)
            )
        else:
            ws.row_dimensions[row_idx].height = 60
